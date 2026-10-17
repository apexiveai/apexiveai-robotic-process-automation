
import os
import json
import secrets
from collections import defaultdict
from datetime import datetime, timezone, timedelta
from contextlib import asynccontextmanager
from io import BytesIO
from typing import Literal

from dotenv import load_dotenv
from fastapi import FastAPI, Depends, HTTPException, Header
from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Font, PatternFill
from pydantic import BaseModel, Field, model_validator
from sqlalchemy import (
    create_engine, String, Integer, Float, Boolean,
    DateTime, Text, select, UniqueConstraint
)
from sqlalchemy.orm import (
    DeclarativeBase, Mapped, mapped_column,
    sessionmaker, Session
)

load_dotenv()

APP_NAME = os.getenv("APP_NAME", "Apexive AI Telecom RPA")
DATABASE_URL = os.getenv(
    "DATABASE_URL", "sqlite:///./apexive_telecom.db"
)
API_TOKEN = os.getenv("API_TOKEN", "")
LICENSE_CODE = os.getenv("LICENSE_CODE", "APEXIVE-DEMO")

if not API_TOKEN or API_TOKEN == "replace-with-a-long-random-secret":
    if os.getenv("APP_ENV", "development") == "production":
        raise RuntimeError("Set a secure API_TOKEN before production startup.")

engine_args = {"pool_pre_ping": True}
if DATABASE_URL.startswith("sqlite"):
    engine_args["connect_args"] = {"check_same_thread": False}

engine = create_engine(DATABASE_URL, **engine_args)
SessionLocal = sessionmaker(bind=engine, expire_on_commit=False)


class Base(DeclarativeBase):
    pass


class KPIRecord(Base):
    __tablename__ = "kpi_records"

    id: Mapped[int] = mapped_column(primary_key=True)
    network_type: Mapped[str] = mapped_column(String(10), index=True)
    metric_name: Mapped[str] = mapped_column(String(100), index=True)
    value: Mapped[float] = mapped_column(Float)
    unit: Mapped[str] = mapped_column(String(30), default="count")
    source: Mapped[str] = mapped_column(String(100), default="sample")
    recorded_at: Mapped[datetime] = mapped_column(DateTime(timezone=True))
    created_at: Mapped[datetime] = mapped_column(
        DateTime(timezone=True),
        default=lambda: datetime.now(timezone.utc)
    )


class KPIThreshold(Base):
    __tablename__ = "kpi_thresholds"
    __table_args__ = (
        UniqueConstraint(
            "network_type", "metric_name",
            name="uq_kpi_threshold_network_metric"
        ),
    )

    id: Mapped[int] = mapped_column(primary_key=True)
    network_type: Mapped[str] = mapped_column(String(10), index=True)
    metric_name: Mapped[str] = mapped_column(String(100), index=True)

    warning_below: Mapped[float | None] = mapped_column(
        Float, nullable=True
    )
    critical_below: Mapped[float | None] = mapped_column(
        Float, nullable=True
    )
    warning_above: Mapped[float | None] = mapped_column(
        Float, nullable=True
    )
    critical_above: Mapped[float | None] = mapped_column(
        Float, nullable=True
    )

    unit: Mapped[str] = mapped_column(String(30), default="percent")
    enabled: Mapped[bool] = mapped_column(Boolean, default=True)
    updated_at: Mapped[datetime] = mapped_column(
        DateTime(timezone=True),
        default=lambda: datetime.now(timezone.utc),
        onupdate=lambda: datetime.now(timezone.utc)
    )


class AnomalyAlert(Base):
    __tablename__ = "anomaly_alerts"

    id: Mapped[int] = mapped_column(primary_key=True)
    kpi_id: Mapped[int] = mapped_column(Integer, index=True)
    network_type: Mapped[str] = mapped_column(String(10), index=True)
    metric_name: Mapped[str] = mapped_column(String(100), index=True)
    observed_value: Mapped[float] = mapped_column(Float)
    severity: Mapped[str] = mapped_column(String(20), index=True)
    message: Mapped[str] = mapped_column(Text)
    status: Mapped[str] = mapped_column(String(20), default="OPEN")
    created_at: Mapped[datetime] = mapped_column(
        DateTime(timezone=True),
        default=lambda: datetime.now(timezone.utc)
    )


class AutomationJob(Base):
    __tablename__ = "automation_jobs"

    id: Mapped[int] = mapped_column(primary_key=True)
    name: Mapped[str] = mapped_column(String(150))
    job_type: Mapped[str] = mapped_column(String(50))
    status: Mapped[str] = mapped_column(String(30), default="QUEUED")
    scheduled_at: Mapped[datetime] = mapped_column(DateTime(timezone=True))
    created_at: Mapped[datetime] = mapped_column(
        DateTime(timezone=True),
        default=lambda: datetime.now(timezone.utc)
    )
    result: Mapped[str | None] = mapped_column(Text, nullable=True)


class AuditLog(Base):
    __tablename__ = "audit_logs"

    id: Mapped[int] = mapped_column(primary_key=True)
    action: Mapped[str] = mapped_column(String(100))
    detail: Mapped[str] = mapped_column(Text)
    created_at: Mapped[datetime] = mapped_column(
        DateTime(timezone=True),
        default=lambda: datetime.now(timezone.utc)
    )


class EnterpriseLicense(Base):
    __tablename__ = "enterprise_licenses"

    id: Mapped[int] = mapped_column(primary_key=True)
    license_code: Mapped[str] = mapped_column(String(150), unique=True)
    customer_name: Mapped[str] = mapped_column(String(150))
    edition: Mapped[str] = mapped_column(String(50), default="ENTERPRISE")
    max_runtimes: Mapped[int] = mapped_column(Integer, default=2)
    active: Mapped[bool] = mapped_column(Boolean, default=True)
    expires_at: Mapped[datetime] = mapped_column(DateTime(timezone=True))


@asynccontextmanager
async def lifespan(app: FastAPI):
    Base.metadata.create_all(engine)

    with SessionLocal() as db:
        license_row = db.scalar(
            select(EnterpriseLicense).where(
                EnterpriseLicense.license_code == LICENSE_CODE
            )
        )

        if not license_row:
            db.add(EnterpriseLicense(
                license_code=LICENSE_CODE,
                customer_name="DEMO CUSTOMER",
                edition="ENTERPRISE",
                max_runtimes=2,
                active=True,
                expires_at=datetime.now(timezone.utc)
                + timedelta(days=30),
            ))

        # Illustrative demo threshold, not an MPT-approved SLA.
        threshold = db.scalar(
            select(KPIThreshold).where(
                KPIThreshold.network_type == "VoLTE",
                KPIThreshold.metric_name == "call_setup_success_rate",
            )
        )

        if not threshold:
            db.add(KPIThreshold(
                network_type="VoLTE",
                metric_name="call_setup_success_rate",
                warning_below=97.0,
                critical_below=95.0,
                unit="percent",
                enabled=True,
            ))

        db.commit()

    from app.hourly_scheduler import start_scheduler, stop_scheduler

    start_scheduler()
    try:
        yield
    finally:
        stop_scheduler()


app = FastAPI(
    title=APP_NAME,
    version="0.2.0",
    description=(
        "Apexive AI Telecom RPA - KPI thresholds and "
        "anomaly alert development API"
    ),
    lifespan=lifespan,
)


def get_db():
    db = SessionLocal()
    try:
        yield db
    finally:
        db.close()


def require_token(
    authorization: str | None = Header(default=None),
):
    expected = f"Bearer {API_TOKEN}"
    if not API_TOKEN or not authorization or not secrets.compare_digest(
        authorization, expected
    ):
        raise HTTPException(
            status_code=401,
            detail="Invalid API credentials"
        )


def audit(db: Session, action: str, detail: str):
    db.add(AuditLog(action=action, detail=detail[:4000]))
    db.commit()


def execute_job_record(db: Session, row: AutomationJob) -> dict:
    """Execute allow-listed, read-only demo workflows."""
    if row.status == "RUNNING":
        raise HTTPException(409, "Job is already running")

    if row.status == "COMPLETED":
        raise HTTPException(409, "Job has already completed")

    scheduled = row.scheduled_at
    if scheduled.tzinfo is None:
        scheduled = scheduled.replace(tzinfo=timezone(timedelta(hours=6, minutes=30)))

    if scheduled > datetime.now(timezone.utc):
        raise HTTPException(409, "Job is scheduled for a future time")

    row.status = "RUNNING"
    db.commit()

    try:
        if row.job_type == "KPI_REPORT":
            report_rows, hourly_items = hourly_report_data(db, 2)
            current_hour = datetime.now(timezone.utc).replace(
                minute=0,
                second=0,
                microsecond=0,
            )
            current_items = [
                item for item in hourly_items if item["hour"] == current_hour
            ]
            network_counts = {
                network: sum(
                    item["samples"] for item in current_items
                    if item["network_type"] == network
                )
                for network in ("CS", "PS", "VoLTE")
            }
            latest_rows = [
                record for record in report_rows
                if (
                    record.recorded_at.replace(tzinfo=timezone.utc)
                    if record.recorded_at.tzinfo is None
                    else record.recorded_at.astimezone(timezone.utc)
                ).replace(minute=0, second=0, microsecond=0) == current_hour
            ]
            latest = max(
                latest_rows,
                key=lambda record: record.recorded_at,
                default=None,
            )

            output = {
                "workflow": "KPI_REPORT",
                "window_hours": 2,
                "records_checked": sum(network_counts.values()),
                "records_by_network": network_counts,
                "current_hour": current_hour.isoformat(),
                "data_status": (
                    "CURRENT_HOUR_DATA_AVAILABLE"
                    if current_items
                    else "NO_CURRENT_HOUR_DATA"
                ),
                "hourly_trends": [
                    {
                        **item,
                        "hour": item["hour"].isoformat(),
                    }
                    for item in current_items
                ],
                "latest_kpi": (
                    {
                        "network_type": latest.network_type,
                        "metric_name": latest.metric_name,
                        "value": latest.value,
                        "unit": latest.unit,
                    }
                    if latest else None
                ),
            }

        elif row.job_type == "KPI_CHECK":
            latest_rows = db.scalars(
                select(KPIRecord)
                .order_by(KPIRecord.recorded_at.desc())
                .limit(1000)
            ).all()

            checked = 0
            for kpi in latest_rows:
                threshold = db.scalar(
                    select(KPIThreshold).where(
                        KPIThreshold.network_type == kpi.network_type,
                        KPIThreshold.metric_name == kpi.metric_name,
                        KPIThreshold.enabled.is_(True),
                    )
                )
                if threshold:
                    checked += 1

            output = {
                "workflow": "KPI_CHECK",
                "records_checked": len(latest_rows),
                "records_with_thresholds": checked,
                "note": "Read-only threshold configuration check",
            }

        elif row.job_type == "ANOMALY_SCAN":
            open_alerts = db.scalars(
                select(AnomalyAlert).where(
                    AnomalyAlert.status == "OPEN"
                )
            ).all()

            critical_count = sum(
                1 for alert in open_alerts
                if alert.severity == "CRITICAL"
            )
            warning_count = sum(
                1 for alert in open_alerts
                if alert.severity == "WARNING"
            )

            output = {
                "workflow": "ANOMALY_SCAN",
                "open_alerts": len(open_alerts),
                "critical_alerts": critical_count,
                "warning_alerts": warning_count,
            }

        else:
            raise ValueError("Unsupported job type")

        row.status = "COMPLETED"
        row.result = json.dumps(output, ensure_ascii=False)
        db.commit()
        db.refresh(row)

        audit(
            db,
            "JOB_EXECUTED",
            f"id={row.id}; type={row.job_type}; status=COMPLETED",
        )

        return {
            "id": row.id,
            "status": row.status,
            "result": output,
        }

    except HTTPException:
        db.rollback()
        raise

    except Exception as exc:
        db.rollback()

        failed_job = db.get(AutomationJob, row.id)
        if failed_job:
            failed_job.status = "FAILED"
            failed_job.result = json.dumps({
                "error": type(exc).__name__,
                "message": "Workflow execution failed",
            })
            db.commit()
            audit(
                db,
                "JOB_EXECUTION_FAILED",
                f"id={row.id}; type={row.job_type}; "
                f"error={type(exc).__name__}",
            )

        raise HTTPException(
            500,
            "Workflow execution failed; inspect audit logs",
        ) from exc


@app.post(
    "/api/v1/jobs/{job_id}/execute",
    dependencies=[Depends(require_token)],
)
def execute_job(
    job_id: int,
    db: Session = Depends(get_db),
):
    row = db.get(AutomationJob, job_id)
    if not row:
        raise HTTPException(404, "Job not found")

    return execute_job_record(db, row)


@app.post(
    "/api/v1/jobs/{job_id}/retry",
    dependencies=[Depends(require_token)],
)
def retry_job(
    job_id: int,
    db: Session = Depends(get_db),
):
    row = db.get(AutomationJob, job_id)
    if not row:
        raise HTTPException(404, "Job not found")

    if row.status not in ("FAILED", "QUEUED"):
        raise HTTPException(
            409,
            "Only FAILED or QUEUED jobs can be retried",
        )

    row.status = "QUEUED"
    row.result = None
    db.commit()

    audit(db, "JOB_REQUEUED", f"id={row.id}; type={row.job_type}")

    # Retry performs one new execution attempt.
    return execute_job_record(db, row)

class KPIInput(BaseModel):
    network_type: Literal["PS", "CS", "VoLTE"]
    metric_name: str = Field(min_length=1, max_length=100)
    value: float
    unit: str = Field(default="count", max_length=30)
    source: str = Field(default="sample", max_length=100)
    recorded_at: datetime


class JobInput(BaseModel):
    name: str = Field(min_length=1, max_length=150)
    job_type: Literal["KPI_REPORT", "KPI_CHECK", "ANOMALY_SCAN"]
    scheduled_at: datetime


class ThresholdInput(BaseModel):
    network_type: Literal["PS", "CS", "VoLTE"]
    metric_name: str = Field(min_length=1, max_length=100)

    warning_below: float | None = None
    critical_below: float | None = None
    warning_above: float | None = None
    critical_above: float | None = None

    unit: str = Field(default="percent", max_length=30)
    enabled: bool = True

    @model_validator(mode="after")
    def validate_thresholds(self):
        if (
            self.warning_below is not None
            and self.critical_below is not None
            and self.critical_below > self.warning_below
        ):
            raise ValueError(
                "critical_below must be <= warning_below"
            )

        if (
            self.warning_above is not None
            and self.critical_above is not None
            and self.critical_above < self.warning_above
        ):
            raise ValueError(
                "critical_above must be >= warning_above"
            )

        if all(value is None for value in (
            self.warning_below,
            self.critical_below,
            self.warning_above,
            self.critical_above,
        )):
            raise ValueError("At least one threshold is required")

        return self


def threshold_severity(value: float, threshold: KPIThreshold) -> str:
    # Critical limits take precedence over warning limits.
    if (
        threshold.critical_below is not None
        and value < threshold.critical_below
    ):
        return "CRITICAL"

    if (
        threshold.critical_above is not None
        and value > threshold.critical_above
    ):
        return "CRITICAL"

    if (
        threshold.warning_below is not None
        and value < threshold.warning_below
    ):
        return "WARNING"

    if (
        threshold.warning_above is not None
        and value > threshold.warning_above
    ):
        return "WARNING"

    return "NORMAL"


def evaluate_kpi(
    db: Session,
    row: KPIRecord,
) -> dict:
    threshold = db.scalar(
        select(KPIThreshold).where(
            KPIThreshold.network_type == row.network_type,
            KPIThreshold.metric_name == row.metric_name,
            KPIThreshold.enabled.is_(True),
        )
    )

    if not threshold:
        return {
            "severity": "NOT_CONFIGURED",
            "alert_id": None,
            "message": "No enabled threshold configured for this KPI",
        }

    severity = threshold_severity(row.value, threshold)

    if severity == "NORMAL":
        return {
            "severity": "NORMAL",
            "alert_id": None,
            "message": "KPI is within configured thresholds",
        }

    message = (
        f"{row.network_type}/{row.metric_name} observed "
        f"{row.value} {row.unit}; severity={severity}"
    )

    alert = AnomalyAlert(
        kpi_id=row.id,
        network_type=row.network_type,
        metric_name=row.metric_name,
        observed_value=row.value,
        severity=severity,
        message=message,
        status="OPEN",
    )
    db.add(alert)
    db.flush()

    db.add(AuditLog(
        action="ANOMALY_ALERT_CREATED",
        detail=(
            f"alert_id={alert.id}; kpi_id={row.id}; "
            f"severity={severity}"
        ),
    ))
    db.commit()
    db.refresh(alert)

    return {
        "severity": severity,
        "alert_id": alert.id,
        "message": message,
    }


def threshold_dict(row: KPIThreshold) -> dict:
    return {
        "id": row.id,
        "network_type": row.network_type,
        "metric_name": row.metric_name,
        "warning_below": row.warning_below,
        "critical_below": row.critical_below,
        "warning_above": row.warning_above,
        "critical_above": row.critical_above,
        "unit": row.unit,
        "enabled": row.enabled,
    }


def alert_dict(row: AnomalyAlert) -> dict:
    return {
        "id": row.id,
        "kpi_id": row.kpi_id,
        "network_type": row.network_type,
        "metric_name": row.metric_name,
        "observed_value": row.observed_value,
        "severity": row.severity,
        "message": row.message,
        "status": row.status,
        "created_at": row.created_at,
    }


def hourly_report_data(
    db: Session,
    hours: int,
    network_type: str | None = None,
) -> tuple[list[KPIRecord], list[dict]]:
    now = datetime.now(timezone.utc)
    cutoff = now - timedelta(hours=hours)
    stmt = select(KPIRecord).where(
        KPIRecord.recorded_at >= cutoff,
        KPIRecord.recorded_at <= now,
    ).order_by(KPIRecord.recorded_at.asc())

    if network_type:
        stmt = stmt.where(KPIRecord.network_type == network_type)

    records = db.scalars(stmt).all()
    buckets: dict[tuple[str, str, str, datetime], list[float]] = defaultdict(list)

    for record in records:
        recorded_at = record.recorded_at
        if recorded_at.tzinfo is None:
            recorded_at = recorded_at.replace(tzinfo=timezone.utc)
        recorded_at = recorded_at.astimezone(timezone.utc)
        hour = recorded_at.replace(minute=0, second=0, microsecond=0)
        buckets[
            (record.network_type, record.metric_name, record.unit, hour)
        ].append(record.value)

    items: list[dict] = []
    previous_values: dict[tuple[str, str, str], float] = {}
    for network, metric, unit, hour in sorted(buckets):
        values = buckets[(network, metric, unit, hour)]
        average = sum(values) / len(values)
        identity = (network, metric, unit)
        previous = previous_values.get(identity)
        change_percent = (
            ((average - previous) / abs(previous)) * 100
            if previous not in (None, 0)
            else None
        )
        direction = (
            "NO_BASELINE" if previous is None
            else "UP" if average > previous
            else "DOWN" if average < previous
            else "FLAT"
        )
        items.append({
            "network_type": network,
            "metric_name": metric,
            "unit": unit,
            "hour": hour,
            "samples": len(values),
            "average": average,
            "minimum": min(values),
            "maximum": max(values),
            "change_percent": change_percent,
            "direction": direction,
        })
        previous_values[identity] = average

    return records, items


def create_hourly_workbook(
    records: list[KPIRecord],
    items: list[dict],
) -> BytesIO:
    workbook = Workbook()
    data_sheet = workbook.active
    data_sheet.title = "KPI Data"
    data_sheet.append([
        "ID", "Network", "Metric", "Value", "Unit", "Source", "Recorded At (UTC)"
    ])

    for record in records:
        recorded_at = record.recorded_at
        if recorded_at.tzinfo is None:
            recorded_at = recorded_at.replace(tzinfo=timezone.utc)
        data_sheet.append([
            record.id,
            record.network_type,
            record.metric_name,
            record.value,
            record.unit,
            record.source,
            recorded_at.astimezone(timezone.utc).replace(tzinfo=None),
        ])

    pivot_sheet = workbook.create_sheet("Hourly Pivot")
    series = sorted({
        (item["network_type"], item["metric_name"], item["unit"])
        for item in items
    })
    hours = sorted({item["hour"] for item in items})
    series_labels = [
        f"{network}/{metric} ({unit})"
        for network, metric, unit in series
    ]
    pivot_sheet.append(["Hour (UTC)", *series_labels])
    values_by_hour = {
        (item["hour"], item["network_type"], item["metric_name"], item["unit"]):
        item["average"]
        for item in items
    }

    for hour in hours:
        pivot_sheet.append([
            hour.replace(tzinfo=None),
            *[
                values_by_hour.get((hour, network, metric, unit))
                for network, metric, unit in series
            ],
        ])

    if hours and series:
        chart = LineChart()
        chart.title = "Hourly KPI Averages by Network and Metric"
        chart.style = 13
        chart.y_axis.title = "Average value"
        chart.x_axis.title = "Hour (UTC)"
        chart.height = 10
        chart.width = 24
        chart.add_data(
            Reference(
                pivot_sheet,
                min_col=2,
                max_col=len(series) + 1,
                min_row=1,
                max_row=len(hours) + 1,
            ),
            titles_from_data=True,
        )
        chart.set_categories(
            Reference(pivot_sheet, min_col=1, min_row=2, max_row=len(hours) + 1)
        )
        pivot_sheet.add_chart(chart, f"A{len(hours) + 4}")

    detail_sheet = workbook.create_sheet("Hourly Detail")
    detail_sheet.append([
        "Network", "Metric", "Unit", "Hour (UTC)", "Samples",
        "Average", "Minimum", "Maximum", "Change (%)", "Direction",
    ])
    for item in items:
        detail_sheet.append([
            item["network_type"],
            item["metric_name"],
            item["unit"],
            item["hour"].replace(tzinfo=None),
            item["samples"],
            item["average"],
            item["minimum"],
            item["maximum"],
            item["change_percent"],
            item["direction"],
        ])

    for sheet in workbook.worksheets:
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="17365D")
        for column in sheet.columns:
            width = min(
                max(
                    len(str(cell.value or ""))
                    for cell in list(column)[:100]
                ) + 2,
                36,
            )
            sheet.column_dimensions[column[0].column_letter].width = width

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output


@app.get("/")
def root():
    return {
        "product": APP_NAME,
        "version": "0.2.0",
        "status": "running",
        "environment": os.getenv("APP_ENV", "development"),
    }


@app.get("/health")
def health():
    return {
        "status": "healthy",
        "timestamp": datetime.now(timezone.utc),
    }


@app.get("/api/v1/license", dependencies=[Depends(require_token)])
def license_status(db: Session = Depends(get_db)):
    row = db.scalar(
        select(EnterpriseLicense).where(
            EnterpriseLicense.license_code == LICENSE_CODE
        )
    )

    if not row:
        raise HTTPException(404, "License not found")

    expiry = row.expires_at
    if expiry.tzinfo is None:
        expiry = expiry.replace(tzinfo=timezone.utc)

    valid = row.active and expiry > datetime.now(timezone.utc)

    return {
        "customer": row.customer_name,
        "edition": row.edition,
        "active": row.active,
        "valid": valid,
        "expires_at": expiry,
        "max_runtimes": row.max_runtimes,
        "mode": "DEVELOPMENT_ENTITLEMENT",
    }


@app.post("/api/v1/kpis", dependencies=[Depends(require_token)])
def create_kpi(data: KPIInput, db: Session = Depends(get_db)):
    values = data.model_dump()
    recorded_at = data.recorded_at
    if recorded_at.tzinfo is None:
        recorded_at = recorded_at.replace(tzinfo=timezone.utc)
    values["recorded_at"] = recorded_at.astimezone(timezone.utc)
    row = KPIRecord(**values)
    db.add(row)
    db.commit()
    db.refresh(row)

    result = evaluate_kpi(db, row)

    db.add(AuditLog(
        action="KPI_INGESTED",
        detail=f"id={row.id}; type={row.network_type}",
    ))
    db.commit()

    return {
        "id": row.id,
        "status": "stored",
        "analysis": result,
    }


@app.get("/api/v1/kpis", dependencies=[Depends(require_token)])
def list_kpis(
    network_type: Literal["PS", "CS", "VoLTE"] | None = None,
    limit: int = 100,
    db: Session = Depends(get_db),
):
    limit = max(1, min(limit, 1000))
    stmt = select(KPIRecord).order_by(
        KPIRecord.recorded_at.desc()
    ).limit(limit)

    if network_type:
        stmt = select(KPIRecord).where(
            KPIRecord.network_type == network_type
        ).order_by(KPIRecord.recorded_at.desc()).limit(limit)

    rows = db.scalars(stmt).all()

    return [
        {
            "id": r.id,
            "network_type": r.network_type,
            "metric_name": r.metric_name,
            "value": r.value,
            "unit": r.unit,
            "source": r.source,
            "recorded_at": r.recorded_at,
        }
        for r in rows
    ]


@app.get("/api/v1/reports/hourly", dependencies=[Depends(require_token)])
def get_hourly_report(
    hours: int = 24,
    network_type: Literal["PS", "CS", "VoLTE"] | None = None,
    db: Session = Depends(get_db),
):
    if not 1 <= hours <= 168:
        raise HTTPException(422, "hours must be between 1 and 168")

    records, items = hourly_report_data(db, hours, network_type)
    return {
        "generated_at": datetime.now(timezone.utc),
        "window_hours": hours,
        "record_count": len(records),
        "items": items,
    }


@app.get(
    "/api/v1/reports/hourly.xlsx",
    dependencies=[Depends(require_token)],
)
def download_hourly_report(
    hours: int = 24,
    network_type: Literal["PS", "CS", "VoLTE"] | None = None,
    db: Session = Depends(get_db),
):
    if not 1 <= hours <= 168:
        raise HTTPException(422, "hours must be between 1 and 168")

    records, items = hourly_report_data(db, hours, network_type)
    workbook = create_hourly_workbook(records, items)
    return Response(
        content=workbook.getvalue(),
        media_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition": (
                'attachment; filename="telecom-hourly-report.xlsx"'
            ),
        },
    )


@app.post("/api/v1/thresholds", dependencies=[Depends(require_token)])
def upsert_threshold(
    data: ThresholdInput,
    db: Session = Depends(get_db),
):
    row = db.scalar(
        select(KPIThreshold).where(
            KPIThreshold.network_type == data.network_type,
            KPIThreshold.metric_name == data.metric_name,
        )
    )

    values = data.model_dump()

    if row:
        for key, value in values.items():
            setattr(row, key, value)
        action = "KPI_THRESHOLD_UPDATED"
    else:
        row = KPIThreshold(**values)
        db.add(row)
        action = "KPI_THRESHOLD_CREATED"

    db.commit()
    db.refresh(row)

    audit(
        db,
        action,
        f"id={row.id}; type={row.network_type}; metric={row.metric_name}",
    )

    return threshold_dict(row)


@app.get("/api/v1/thresholds", dependencies=[Depends(require_token)])
def list_thresholds(db: Session = Depends(get_db)):
    rows = db.scalars(
        select(KPIThreshold).order_by(
            KPIThreshold.network_type,
            KPIThreshold.metric_name,
        )
    ).all()
    return [threshold_dict(row) for row in rows]


@app.get("/api/v1/alerts", dependencies=[Depends(require_token)])
def list_alerts(
    severity: Literal["WARNING", "CRITICAL"] | None = None,
    status: Literal["OPEN", "RESOLVED"] | None = None,
    limit: int = 100,
    db: Session = Depends(get_db),
):
    limit = max(1, min(limit, 1000))
    stmt = select(AnomalyAlert).order_by(
        AnomalyAlert.id.desc()
    )

    if severity:
        stmt = stmt.where(AnomalyAlert.severity == severity)

    if status:
        stmt = stmt.where(AnomalyAlert.status == status)

    rows = db.scalars(stmt.limit(limit)).all()
    return [alert_dict(row) for row in rows]


@app.post(
    "/api/v1/alerts/{alert_id}/resolve",
    dependencies=[Depends(require_token)],
)
def resolve_alert(
    alert_id: int,
    db: Session = Depends(get_db),
):
    row = db.get(AnomalyAlert, alert_id)
    if not row:
        raise HTTPException(404, "Alert not found")

    row.status = "RESOLVED"
    db.commit()
    db.refresh(row)

    audit(db, "ANOMALY_ALERT_RESOLVED", f"alert_id={row.id}")
    return alert_dict(row)


@app.post("/api/v1/jobs", dependencies=[Depends(require_token)])
def create_job(data: JobInput, db: Session = Depends(get_db)):
    if data.scheduled_at.tzinfo is None:
        raise HTTPException(
            422,
            "scheduled_at must include a timezone, e.g. +06:30",
        )

    if data.scheduled_at < datetime.now(timezone.utc):
        raise HTTPException(422, "scheduled_at cannot be in the past")

    row = AutomationJob(
        name=data.name,
        job_type=data.job_type,
        scheduled_at=data.scheduled_at,
        status="QUEUED",
    )
    db.add(row)
    db.commit()
    db.refresh(row)

    audit(db, "JOB_CREATED", f"id={row.id}; type={row.job_type}")
    return {"id": row.id, "status": row.status}


@app.get("/api/v1/jobs", dependencies=[Depends(require_token)])
def list_jobs(db: Session = Depends(get_db)):
    rows = db.scalars(
        select(AutomationJob).order_by(
            AutomationJob.id.desc()
        ).limit(200)
    ).all()

    return [
        {
            "id": r.id,
            "name": r.name,
            "job_type": r.job_type,
            "status": r.status,
            "scheduled_at": r.scheduled_at,
            "result": r.result,
        }
        for r in rows
    ]


@app.get("/api/v1/audit-logs", dependencies=[Depends(require_token)])
def list_audit_logs(db: Session = Depends(get_db)):
    rows = db.scalars(
        select(AuditLog).order_by(
            AuditLog.id.desc()
        ).limit(200)
    ).all()

    return [
        {
            "id": r.id,
            "action": r.action,
            "detail": r.detail,
            "created_at": r.created_at,
        }
        for r in rows
    ]
